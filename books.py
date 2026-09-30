"""The bookkeeping rules: sheet layouts and every calculated figure.

Purchases, Sales and Expenses are entered by hand (through the app, or
directly in the spreadsheet). Everything else — totals, landed costs, profit,
the Stock and Monthly Summary sheets — is recalculated on every save.

Each row is recorded in its own currency (INR or USD) together with the
exchange rate on that day (₹ for $1), and the ₹ and $ equivalents are worked
out from that rate, so the books can be read in either currency.

An expense can optionally name one or more Saree Codes. It is then added to
the cost of those sarees (spread over the pieces bought) instead of being
counted as a general business expense.
"""
from __future__ import annotations

import io
from datetime import datetime, timedelta

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CURRENCIES = ["INR", "USD"]
RATE = "Rate (₹ per $1)"
LINKED = "Added to Saree Cost"

PURCHASE_COLS = [
    "Date", "Supplier", "Invoice No", "Saree Code", "Description", "Fabric / Type",
    "Colour", "Qty", "Pricing", "Currency", RATE, "Cost per Piece", "Shipping + Customs",
    "Total Cost", "Landed Cost per Piece", "Total Cost ₹", "Total Cost $", "Notes",
]
SALE_COLS = [
    "Date", "Bill No", "Customer", "Phone", "Saree Code", "Description", "Qty",
    "Currency", RATE, "Price per Piece", "Discount", "Total Amount", "Amount Paid",
    "Balance Due", "Payment Mode", "Total ₹", "Total $", "Balance Due ₹", "Balance Due $",
    "Cost of Goods ₹", "Cost of Goods $", "Profit ₹", "Profit $", "Notes",
]
EXPENSE_COLS = [
    "Date", "Category", "Description", "Saree Code", "Currency", RATE, "Amount",
    "Amount ₹", "Amount $", LINKED, "Payment Mode", "Notes",
]
STOCK_COLS = [
    "Saree Code", "Description", "Fabric / Type", "Qty Bought", "Qty Sold", "In Stock",
    "Extra Costs ₹", "Extra Costs $", "Avg Landed Cost ₹", "Avg Landed Cost $",
    "Stock Value ₹", "Stock Value $",
]
SUMMARY_METRICS = [
    "Purchases", "Sales", "Cost of Goods Sold", "Gross Profit", "Expenses", "Net Profit",
    "Pending Payments",
]
SUMMARY_COLS = ["Month"] + [f"{m} {c}" for m in SUMMARY_METRICS for c in ("₹", "$")]

INPUT_SHEETS = {"Purchases": PURCHASE_COLS, "Sales": SALE_COLS, "Expenses": EXPENSE_COLS}
ALL_SHEETS = list(INPUT_SHEETS) + ["Stock", "Monthly Summary"]
DEFAULT_CURRENCY = {"Purchases": "INR", "Sales": "USD", "Expenses": "INR"}

# Columns the app fills in itself — greyed out when editing in the app.
COMPUTED = {
    "Purchases": ["Total Cost", "Landed Cost per Piece", "Total Cost ₹", "Total Cost $"],
    "Sales": ["Total Amount", "Balance Due", "Total ₹", "Total $", "Balance Due ₹",
              "Balance Due $", "Cost of Goods ₹", "Cost of Goods $", "Profit ₹", "Profit $"],
    "Expenses": ["Amount ₹", "Amount $", LINKED],
}

QTY_COLS = {"Qty", "Qty Bought", "Qty Sold", "In Stock"}
TEXT_COLS = {
    "Date", "Month", "Supplier", "Invoice No", "Saree Code", "Description", "Fabric / Type",
    "Colour", "Pricing", "Currency", "Bill No", "Customer", "Phone", "Payment Mode",
    "Category", "Notes", LINKED,
}


def is_number_col(col: str) -> bool:
    return col not in TEXT_COLS


def _empty(cols: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=cols)


def _as_text(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _as_date(v):
    """Dates may arrive as date objects, text, or spreadsheet serial numbers."""
    if isinstance(v, (int, float)) and not pd.isna(v) and v > 0:
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).date()
    d = pd.to_datetime(v, errors="coerce", dayfirst=isinstance(v, str) and "/" in v)
    return None if pd.isna(d) else d.date()


def clean(df: pd.DataFrame, cols: list[str], default_currency: str) -> pd.DataFrame:
    """Ensure the expected columns exist, in order, with sensible types."""
    df = df.copy()
    for c in cols:
        if c not in df.columns:
            df[c] = None
    df = df[cols].replace("", None).dropna(how="all")
    df["Date"] = df["Date"].apply(_as_date)
    for c in cols:
        if is_number_col(c):
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
        elif c != "Date":
            df[c] = df[c].apply(_as_text)
    cur = df["Currency"].str.upper()
    df["Currency"] = cur.where(cur.isin(CURRENCIES), default_currency)
    return df.reset_index(drop=True)


def clean_all(raw: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    return {
        name: clean(raw.get(name, _empty(cols)), cols, DEFAULT_CURRENCY[name])
        for name, cols in INPUT_SHEETS.items()
    }


# ---------- currency ----------

def convert(amount, currency: str, rate: float, to: str):
    """Convert between INR and USD using a rate expressed as ₹ for $1."""
    if currency == to:
        return amount
    if not rate:
        return 0
    return amount * rate if to == "INR" else amount / rate


def _both(df: pd.DataFrame, src: str, inr_col: str, usd_col: str) -> None:
    rate = df[RATE].where(df[RATE] > 0)
    is_inr = df["Currency"] == "INR"
    df[inr_col] = df[src].where(is_inr, df[src] * rate).fillna(0).round(2)
    df[usd_col] = df[src].where(~is_inr, df[src] / rate).fillna(0).round(2)


def split_codes(text: str) -> list[str]:
    return [c.strip().upper() for c in str(text).split(",") if c.strip()]


# ---------- derived figures ----------

def recalc_purchases(p: pd.DataFrame) -> pd.DataFrame:
    p = p.copy()
    p["Total Cost"] = (p["Qty"] * p["Cost per Piece"] + p["Shipping + Customs"]).round(2)
    p["Landed Cost per Piece"] = (p["Total Cost"] / p["Qty"].where(p["Qty"] != 0)).fillna(0).round(2)
    _both(p, "Total Cost", "Total Cost ₹", "Total Cost $")
    return p


def _qty_bought(p: pd.DataFrame) -> pd.Series:
    return p.groupby("Saree Code")["Qty"].sum() if not p.empty else pd.Series(dtype=float)


def recalc_expenses(e: pd.DataFrame, p: pd.DataFrame) -> pd.DataFrame:
    e = e.copy()
    _both(e, "Amount", "Amount ₹", "Amount $")
    bought = _qty_bought(p)
    e["Saree Code"] = e["Saree Code"].apply(lambda t: ", ".join(split_codes(t)))
    e[LINKED] = e["Saree Code"].apply(
        lambda t: "Yes" if any(bought.get(c, 0) > 0 for c in split_codes(t)) else ""
    )
    return e


def extra_costs(p: pd.DataFrame, e: pd.DataFrame) -> pd.DataFrame:
    """Saree-linked expenses per design, spread by pieces bought when several are named."""
    bought = _qty_bought(p)
    out: dict[str, list[float]] = {}
    for _, r in e[e[LINKED] == "Yes"].iterrows():
        codes = [c for c in split_codes(r["Saree Code"]) if bought.get(c, 0) > 0]
        total = sum(bought[c] for c in codes)
        for c in codes:
            share = bought[c] / total
            acc = out.setdefault(c, [0.0, 0.0])
            acc[0] += r["Amount ₹"] * share
            acc[1] += r["Amount $"] * share
    return pd.DataFrame.from_dict(out, orient="index", columns=["₹", "$"])


def avg_landed_cost(p: pd.DataFrame, e: pd.DataFrame) -> pd.DataFrame:
    """Average landed cost per piece of each design (incl. linked expenses), in ₹ and $."""
    if p.empty:
        return pd.DataFrame(columns=["₹", "$"])
    g = p.groupby("Saree Code")[["Total Cost ₹", "Total Cost $", "Qty"]].sum()
    extra = extra_costs(p, e).reindex(g.index).fillna(0)
    qty = g["Qty"].where(g["Qty"] != 0)
    return pd.DataFrame({
        "₹": ((g["Total Cost ₹"] + extra["₹"]) / qty).fillna(0),
        "$": ((g["Total Cost $"] + extra["$"]) / qty).fillna(0),
    })


def recalc_sales(s: pd.DataFrame, p: pd.DataFrame, e: pd.DataFrame) -> pd.DataFrame:
    s = s.copy()
    cost = avg_landed_cost(p, e)
    s["Total Amount"] = (s["Qty"] * s["Price per Piece"] - s["Discount"]).round(2)
    s["Balance Due"] = (s["Total Amount"] - s["Amount Paid"]).round(2)
    _both(s, "Total Amount", "Total ₹", "Total $")
    _both(s, "Balance Due", "Balance Due ₹", "Balance Due $")
    for c in ("₹", "$"):
        s[f"Cost of Goods {c}"] = (s["Qty"] * s["Saree Code"].map(cost[c]).fillna(0)).round(2)
        s[f"Profit {c}"] = (s[f"Total {c}"] - s[f"Cost of Goods {c}"]).round(2)
    return s


def build_stock(p: pd.DataFrame, s: pd.DataFrame, e: pd.DataFrame) -> pd.DataFrame:
    if p.empty:
        return _empty(STOCK_COLS)
    info = p.groupby("Saree Code").agg(
        **{"Description": ("Description", "last"), "Fabric / Type": ("Fabric / Type", "last"),
           "Qty Bought": ("Qty", "sum")}
    )
    sold = s.groupby("Saree Code")["Qty"].sum() if not s.empty else pd.Series(dtype=float)
    info["Qty Sold"] = sold.reindex(info.index).fillna(0)
    info["In Stock"] = info["Qty Bought"] - info["Qty Sold"]
    extra = extra_costs(p, e).reindex(info.index).fillna(0)
    cost = avg_landed_cost(p, e)
    for c in ("₹", "$"):
        info[f"Extra Costs {c}"] = extra[c].round(2)
        info[f"Avg Landed Cost {c}"] = cost[c].reindex(info.index).fillna(0).round(2)
        info[f"Stock Value {c}"] = (info["In Stock"] * info[f"Avg Landed Cost {c}"]).round(2)
    return info.reset_index()[STOCK_COLS].sort_values("Saree Code")


def general_expenses(e: pd.DataFrame) -> pd.DataFrame:
    """Expenses not already added to the cost of particular sarees."""
    return e[e[LINKED] != "Yes"]


def build_summary(p: pd.DataFrame, s: pd.DataFrame, e: pd.DataFrame) -> pd.DataFrame:
    def by_month(df: pd.DataFrame, col: str) -> pd.Series:
        if df.empty:
            return pd.Series(dtype=float)
        m = pd.to_datetime(df["Date"], errors="coerce").dt.strftime("%Y-%m")
        return df.groupby(m)[col].sum()

    linked, general = e[e[LINKED] == "Yes"], general_expenses(e)
    cols = {}
    for c in ("₹", "$"):
        cols[f"Purchases {c}"] = by_month(p, f"Total Cost {c}").add(by_month(linked, f"Amount {c}"), fill_value=0)
        cols[f"Sales {c}"] = by_month(s, f"Total {c}")
        cols[f"Cost of Goods Sold {c}"] = by_month(s, f"Cost of Goods {c}")
        cols[f"Expenses {c}"] = by_month(general, f"Amount {c}")
        cols[f"Pending Payments {c}"] = by_month(s, f"Balance Due {c}")
    out = pd.DataFrame(cols).fillna(0)
    if out.empty:
        return _empty(SUMMARY_COLS)
    for c in ("₹", "$"):
        out[f"Gross Profit {c}"] = out[f"Sales {c}"] - out[f"Cost of Goods Sold {c}"]
        out[f"Net Profit {c}"] = out[f"Gross Profit {c}"] - out[f"Expenses {c}"]
    out = out.sort_index().reset_index(names="Month")
    return out[SUMMARY_COLS].round(2)


def compute(data: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """Clean the entered sheets and build every sheet of the workbook."""
    data = clean_all(data)
    p = recalc_purchases(data["Purchases"])
    e = recalc_expenses(data["Expenses"], p)
    s = recalc_sales(data["Sales"], p, e)
    by_date = dict(kind="stable", na_position="last")
    return {
        "Purchases": p.sort_values("Date", **by_date).reset_index(drop=True),
        "Sales": s.sort_values("Date", **by_date).reset_index(drop=True),
        "Expenses": e.sort_values("Date", **by_date).reset_index(drop=True),
        "Stock": build_stock(p, s, e).reset_index(drop=True),
        "Monthly Summary": build_summary(p, s, e),
    }


# ---------- Excel output ----------

HEADER_FILL = PatternFill("solid", fgColor="7B1E3A")  # deep maroon
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _format_sheet(ws, df: pd.DataFrame) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, start=1):
        letter = get_column_letter(i)
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()[:500]])
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 40)
        if col == "Date":
            fmt = "DD-MMM-YYYY"
        elif col in QTY_COLS:
            fmt = "0"
        elif is_number_col(col):
            fmt = "#,##0.00"
        else:
            continue
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            cell.number_format = fmt
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def write_xlsx(sheets: dict[str, pd.DataFrame], target) -> None:
    """Write all sheets, formatted, to a path or file-like object."""
    with pd.ExcelWriter(target, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            _format_sheet(xw.sheets[name], df)


def xlsx_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    buf = io.BytesIO()
    write_xlsx(sheets, buf)
    return buf.getvalue()


# ---------- helpers for the entry forms ----------

def last_rate(data: dict[str, pd.DataFrame]) -> float | None:
    """Most recent exchange rate used anywhere — a fallback when offline."""
    rows = pd.concat([df[["Date", RATE]] for df in data.values() if RATE in df])
    rows = rows[rows[RATE] > 0].sort_values("Date")
    return float(rows[RATE].iloc[-1]) if not rows.empty else None


def next_bill_no(sales: pd.DataFrame) -> str:
    nums = pd.to_numeric(sales["Bill No"].str.extract(r"(\d+)$")[0], errors="coerce").dropna()
    return f"B{int(nums.max()) + 1 if not nums.empty else 1:04d}"


def split_lot(items: pd.DataFrame, bulk_price: float, shipping: float) -> pd.DataFrame:
    """Work out cost per piece for a purchase lot.

    `items` has Qty and an optional "Own Price" per piece. Pieces without their
    own price share `bulk_price` equally. `shipping` (already in the lot's
    currency) is spread equally over every piece in the lot.
    """
    items = items.copy()
    own = pd.to_numeric(items["Own Price"], errors="coerce")
    has_own = own.notna() & (own > 0)
    bulk_qty = items.loc[~has_own, "Qty"].sum()
    bulk_each = bulk_price / bulk_qty if bulk_qty else 0
    total_qty = items["Qty"].sum()
    ship_each = shipping / total_qty if total_qty else 0
    items["Pricing"] = has_own.map({True: "Own price", False: "Bulk lot share"})
    items["Cost per Piece"] = own.where(has_own, bulk_each).round(2)
    items["Shipping + Customs"] = (ship_each * items["Qty"]).round(2)
    items["Landed Cost per Piece"] = (items["Cost per Piece"] + ship_each).round(2)
    return items

